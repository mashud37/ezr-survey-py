"""Build PowerPoint decks and Word documents slide by slide, placing charts and tables by the template's own placeholders. Any organisation template works, whatever its layouts are called."""

import copy
import os
import tempfile
from pathlib import Path

import pandas as pd

from .config import ezrsurvey_default
from .export import default_output_path, is_ggplot, resolve_output_path
from .rbase import character_value, is_missing, match_arg, r_round, r_sort, warn

TEMPLATE_FOLDER = Path(__file__).resolve().parent / "templates"
FORMATS = ["pptx", "docx"]
STYLES = ["elevated", "plain"]
PURPOSES = ["content", "title", "two_content"]
EMU_PER_INCH = 914400
REPORT_STATE = {}
SLIDE_NUMBER_ID = 900
FALLBACK_SLOT = {"left": 0.5, "top": 1.75, "width": 9, "height": 5}
DOCX_PLOT_SIZE = {"width": 6, "height": 3.5}
PLOT_DPI = 150
TABLE_FONT_SIZE = 12
TABLE_PADDING_POINTS = 4
NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
PLACEHOLDER_TYPES = {
    "OBJECT": "body",
    "BODY": "body",
    "TITLE": "title",
    "CENTER_TITLE": "ctrTitle",
    "SUBTITLE": "subTitle",
    "DATE": "dt",
    "FOOTER": "ftr",
    "SLIDE_NUMBER": "sldNum",
    "PICTURE": "pic",
    "TABLE": "tbl",
    "CHART": "chart",
    "MEDIA_CLIP": "media",
    "ORG_CHART": "dgm",
}


def require_pptx():
    try:
        import pptx  # noqa: F401
    except ImportError:
        raise ValueError("Package 'python-pptx' is required for the report builders. Install it with pip install python-pptx.") from None


def require_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        raise ValueError("Package 'python-docx' is required for Word reports. Install it with pip install python-docx.") from None


def is_pptx(doc):
    return type(doc).__module__.startswith("pptx.") and type(doc).__name__ == "Presentation"


def is_docx(doc):
    return type(doc).__module__.startswith("docx.") and type(doc).__name__ == "Document"


def check_doc(doc):
    if not is_pptx(doc) and not is_docx(doc):
        raise ValueError("`doc` must be an ezrsurvey report (see report_new()).")


def report_state(doc):
    return REPORT_STATE.setdefault(id(doc), {})


def set_report_state(doc, changes):
    state = report_state(doc)
    state.update(changes)
    return state


def default_pptx_template(style="elevated"):
    style = match_arg(style, STYLES)
    name = "ezrsurvey-16x9-plain.pptx" if style == "plain" else "ezrsurvey-16x9.pptx"
    path = TEMPLATE_FOLDER / name
    return str(path) if path.exists() else None


def master_name(master):
    """The master's name as officer reports it: the name of the theme the master uses."""
    from pptx.opc.constants import RELATIONSHIP_TYPE

    for relationship in master.part.rels.values():
        if relationship.reltype == RELATIONSHIP_TYPE.THEME:
            text = relationship.target_part.blob.decode("utf-8", errors="ignore")
            start = text.find('name="')
            if start >= 0:
                return text[start + 6 : text.find('"', start + 6)]
    return master.name


def inches(value):
    return None if value is None else value / EMU_PER_INCH


def layout_props(doc):
    """Every placeholder of every layout: master, layout, type, label and geometry in inches."""
    rows = []
    for master in doc.slide_masters:
        master_label = master_name(master)
        for layout in master.slide_layouts:
            for placeholder in layout.placeholders:
                kind = str(placeholder.placeholder_format.type).split(".")[-1].split(" ")[0]
                rows.append(
                    {
                        "master_name": master_label,
                        "name": layout.name,
                        "type": PLACEHOLDER_TYPES.get(kind, kind.lower()),
                        "id": placeholder.shape_id,
                        "ph_label": placeholder.name,
                        "offx": inches(placeholder.left),
                        "offy": inches(placeholder.top),
                        "cx": inches(placeholder.width),
                        "cy": inches(placeholder.height),
                        "layout": layout,
                        "placeholder": placeholder,
                    }
                )
    return rows


def layout_groups(doc):
    """Placeholder rows grouped by layout, in the order R's split() gives them."""
    groups = {}
    for row in layout_props(doc):
        groups.setdefault(row["master_name"] + "\r" + row["name"], []).append(row)
    ordered = {}
    for key in r_sort(groups.keys()):
        ordered[key] = groups[key]
    return ordered


def layout_score(rows, purpose):
    types = [row["type"] for row in rows]
    n_body = types.count("body")
    has_title = "title" in types
    has_ctr = "ctrTitle" in types
    name = rows[0]["name"].lower()
    if purpose == "content":
        named = any(word in name for word in ("content", "body", "inhalt", "contenu"))
        return 2 * has_title + 2 * (n_body == 1) + 1 * (n_body > 1) - 2 * has_ctr + 0.5 * named
    if purpose == "title":
        named = any(word in name for word in ("title", "titel", "titre"))
        return 3 * has_ctr + 0.5 * has_title + 1 * named
    named = any(word in name for word in ("two", "comparison", "zwei", "deux"))
    return 3 * (n_body == 2 and has_title) + 0.5 * named


def select_layout(doc, purpose="content"):
    """The layout best suited to `purpose`, scored by placeholder types rather than by name.

    Returns:
        A dict with ``layout`` and ``master`` names, or None when no layout has
        two content placeholders and a title (``two_content`` only).
    """
    purpose = match_arg(purpose, PURPOSES)
    groups = list(layout_groups(doc).values())
    scores = [layout_score(rows, purpose) for rows in groups]
    best = scores.index(max(scores))
    if purpose == "two_content" and scores[best] < 3:
        return None
    return {"layout": groups[best][0]["name"], "master": groups[best][0]["master_name"]}


def layout_placeholders(doc, layout, master, types=("body",)):
    """The placeholders of `types` on one layout, in reading order (top to bottom, then left to right)."""
    types = [types] if isinstance(types, str) else list(types)
    rows = [row for row in layout_props(doc) if row["name"] == layout and row["master_name"] == master and row["type"] in types]
    return sorted(rows, key=reading_order)


def reading_order(row):
    """Top to bottom, then left to right; placeholders without a position go last."""
    top = float("inf") if row["offy"] is None else row["offy"]
    left = float("inf") if row["offx"] is None else row["offx"]
    return (top, left)


def find_layout(doc, layout, master=None):
    for slide_master in doc.slide_masters:
        for candidate in slide_master.slide_layouts:
            if candidate.name == layout and (master is None or master_name(slide_master) == master):
                return candidate
    return None


def content_slot(doc, index=1):
    """Where, and how large, content should land on the current slide, in inches.

    Falls back from the layout's body placeholder to the slide size minus
    margins, then to a fixed 9 by 5 inches.
    """
    state = report_state(doc)
    rows = []
    if state.get("layout") is not None:
        rows = layout_placeholders(doc, state["layout"], state["master"], "body")
    if len(rows) >= index:
        row = rows[index - 1]
        if row["cx"] and row["cy"]:
            return {"left": row["offx"], "top": row["offy"], "width": row["cx"], "height": row["cy"], "label": row["ph_label"]}
    width = inches(doc.slide_width)
    height = inches(doc.slide_height)
    if width and height:
        return {"left": 0.5, "top": 1.75, "width": width - 1, "height": height - 2.25, "label": None}
    return dict(FALLBACK_SLOT, label=None)


def current_slide(doc):
    return doc.slides[len(doc.slides) - 1]


def add_slide_number(doc, layout, master):
    """Copy the layout's slide-number placeholder onto the current slide, as officer never does."""
    if layout_placeholders(doc, layout, master, "ctrTitle"):
        return doc
    numbers = layout_placeholders(doc, layout, master, "sldNum")
    if not numbers:
        return doc
    slide = current_slide(doc)
    element = copy.deepcopy(numbers[0]["placeholder"]._element)
    identity = element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    if identity is not None:
        identity.set("id", str(SLIDE_NUMBER_ID + len(doc.slides)))
    slide.shapes._spTree.append(element)
    return doc


def report_layouts(template=None, style="elevated"):
    """List the slide layouts of a PowerPoint template.

    Shows what a template offers the report builders: every layout, its master,
    and the placeholders it carries. Useful to decide which `layout` to name in
    `report_add_slide()`, or to sanity-check an organisation template before
    building a deck against it. The report builders pick layouts automatically
    by inspecting placeholders, so most decks never need this.

    Parameters
    ----------
    template : str or pptx.Presentation, optional
        Path to a .pptx, an existing document from `report_new()`, or None
        (default) for the brand template registered by `use_brand()`, falling
        back to a built-in template.
    style : str
        Built-in template to inspect when no template is set: "elevated"
        (default) or "plain".

    Returns
    -------
    pandas.DataFrame
        One row per layout: ``layout``, ``master``, ``has_title``, ``n_body``,
        ``body_width``, ``body_height`` (inches of the first body placeholder,
        missing when the layout has none).

    See Also
    --------
    report_new, report_add_slide, use_brand

    Examples
    --------
    >>> report_layouts()
    """
    require_pptx()
    from pptx import Presentation

    style = match_arg(style, STYLES)
    if is_pptx(template):
        doc = template
    else:
        path = template or ezrsurvey_default("brand_template_pptx") or default_pptx_template(style)
        doc = Presentation(path)
    rows = []
    for placeholders in layout_groups(doc).values():
        bodies = [row for row in placeholders if row["type"] == "body"]
        rows.append(
            {
                "layout": placeholders[0]["name"],
                "master": placeholders[0]["master_name"],
                "has_title": any(row["type"] in ("title", "ctrTitle") for row in placeholders),
                "n_body": len(bodies),
                "body_width": r_round(bodies[0]["cx"], 2) if bodies else float("nan"),
                "body_height": r_round(bodies[0]["cy"], 2) if bodies else float("nan"),
            }
        )
    return pd.DataFrame(rows, columns=["layout", "master", "has_title", "n_body", "body_width", "body_height"])


def remove_all_slides(doc):
    slide_list = doc.slides._sldIdLst
    for slide_id in list(slide_list):
        doc.part.drop_rel(slide_id.rId)
        slide_list.remove(slide_id)


def report_new(format="pptx", template=None, style="elevated", slide_numbers=True, keep_slides=True):  # lint-style: ignore FN001
    """Start a new PowerPoint or Word report.

    Opens a python-pptx or python-docx document you can build up with the
    ``report_add_*()`` helpers and write out with `report_save()`: the "build it
    from Python" path. Slide layouts are chosen by inspecting the template's
    placeholders, so any organisation template works; see `report_layouts()`
    for what yours contains. For a quick deck from a list of charts and tables,
    `report_deck()` does the whole thing in one call; for a document-driven
    workflow, `scaffold_report()` gives you a Quarto template instead. Needs
    the ``reports`` extra.

    Parameters
    ----------
    format : str
        "pptx" (default) for PowerPoint or "docx" for Word.
    template : str, optional
        Path to a .pptx / .docx to use as the style template. None (default)
        uses the brand template registered by `use_brand()`, falling back to one
        of the package's built-in 16:9 templates (see `style`).
    style : str
        Built-in template used when no template is set. "elevated" (default) is
        the styled deck: a navy/gold identity with a full-bleed navy cover,
        navy section dividers, a navy title over one slim rule on every content
        slide and slide numbers in the corner. "plain" is the same widescreen
        deck with no decoration.
    slide_numbers : bool
        If True (default), every content slide added with `report_add_slide()`
        shows the template's slide number.
    keep_slides : bool
        PowerPoint only. If True (default), slides already in the template stay
        and yours are added after them. False starts from an empty deck.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_add_slide, report_add_plot, report_add_table, report_save, report_deck

    Examples
    --------
    >>> doc = report_new("pptx")
    """
    format = match_arg(format, FORMATS)
    style = match_arg(style, STYLES)
    template = template or ezrsurvey_default(f"brand_template_{format}")
    if format == "pptx":
        require_pptx()
        from pptx import Presentation

        doc = Presentation(template or default_pptx_template(style))
        if not keep_slides:
            remove_all_slides(doc)
        set_report_state(doc, {"slide_numbers": bool(slide_numbers), "added": []})
        return doc
    require_docx()
    import docx

    return docx.Document(template) if template else docx.Document()


def chosen_layout(doc, layout, master):
    if layout is None:
        return select_layout(doc, "content")
    listed = report_layouts(doc)
    hits = listed[listed["layout"] == layout]
    if master is not None:
        hits = hits[hits["master"] == master]
    if len(hits) == 0:
        raise ValueError(f"Layout '{layout}' not found in this template. See report_layouts() for what it offers.")
    return {"layout": hits["layout"].iloc[0], "master": hits["master"].iloc[0]}


def fill_placeholder(doc, row, text):
    """Clone one layout placeholder onto the current slide and write text into it."""
    slide = current_slide(doc)
    slide.shapes.clone_placeholder(row["placeholder"])
    placeholder = slide.shapes[len(slide.shapes) - 1]
    lines = [text] if isinstance(text, str) else list(text)
    placeholder.text_frame.text = lines[0]
    for line in lines[1:]:
        placeholder.text_frame.add_paragraph().text = line
    return placeholder


def add_empty_slide(doc, selection):
    """A new slide on the chosen layout, bare as officer makes it: no placeholders until filled."""
    layout = find_layout(doc, selection["layout"], selection["master"])
    slide = doc.slides.add_slide(layout)
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    return slide


def report_add_slide(doc, title=None, layout=None, master=None, heading_level=1):  # lint-style: ignore FN001
    """Add a slide (PowerPoint) or heading (Word) to a report.

    For a pptx document this starts a new slide and sets its title; for a docx
    document it adds a heading paragraph. Automatic selection scores every
    layout of the template: a title placeholder plus a single content
    placeholder is the ideal, and the layout's *placeholders* decide, not its
    (language-dependent) name. If the chosen layout has no title placeholder,
    the title is skipped with a warning rather than failing the build.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    title : str, optional
        Slide title / heading text.
    layout, master : str, optional
        PowerPoint layout and master names. None (default) picks the
        template's best content layout automatically. Name a layout from
        `report_layouts()` to override.
    heading_level : int
        Word heading level (1-3). Default 1.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_new, report_layouts

    Examples
    --------
    >>> doc = report_add_slide(report_new("pptx"), "Audience")
    """
    check_doc(doc)
    if not is_pptx(doc):
        if title is not None:
            doc.add_paragraph(title, style=f"Heading {heading_level}")
        return doc
    selection = chosen_layout(doc, layout, master)
    slide = add_empty_slide(doc, selection)
    state = set_report_state(doc, {"layout": selection["layout"], "master": selection["master"]})
    state.setdefault("added", []).append(slide.slide_id)
    if state.get("slide_numbers") is not False:
        add_slide_number(doc, selection["layout"], selection["master"])
    if title is None:
        return doc
    titles = layout_placeholders(doc, selection["layout"], selection["master"], ["title", "ctrTitle"])
    if titles:
        fill_placeholder(doc, titles[0], title)
    elif not state.get("warned_title"):
        warn(f"Layout '{selection['layout']}' has no title placeholder; slide titles are skipped.")
        set_report_state(doc, {"warned_title": True})
    return doc


def render_png(plot, width, height, dpi):
    handle = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    handle.close()
    plot.save(handle.name, width=width, height=height, dpi=dpi, units="in", verbose=False, facecolor="white")
    return handle.name


def report_add_plot(doc, plot, width=None, height=None, dpi=PLOT_DPI):
    """Add a plot to a report.

    Renders a chart to an image and places it in the current slide's content
    placeholder (pptx) or as a new figure (docx). On slides the image is
    rendered at exactly the content placeholder's size, so nothing is stretched
    or letterboxed; when the template does not define placeholder geometry the
    slide size (minus margins) is used instead.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    plot : plotnine.ggplot
        A chart.
    width, height : float, optional
        Image size in inches. None (default) sizes it to the slide's content
        placeholder (pptx) or 6 x 3.5 (docx).
    dpi : int
        Raster resolution. Default 150.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_add_table

    Examples
    --------
    >>> p = plot_bars(calc_percentage(podracing_survey, "demo_gender"))
    >>> doc = report_add_slide(report_new("pptx"), "Gender")
    >>> doc = report_add_plot(doc, p)
    """
    check_doc(doc)
    if is_pptx(doc):
        from pptx.util import Inches

        slot = content_slot(doc)
        width = slot["width"] if width is None else width
        height = slot["height"] if height is None else height
        image = render_png(plot, width, height, dpi)
        current_slide(doc).shapes.add_picture(image, Inches(slot["left"]), Inches(slot["top"]), Inches(width), Inches(height))
        os.remove(image)
        return doc
    from docx.shared import Inches

    width = DOCX_PLOT_SIZE["width"] if width is None else width
    height = DOCX_PLOT_SIZE["height"] if height is None else height
    image = render_png(plot, width, height, dpi)
    doc.add_picture(image, width=Inches(width), height=Inches(height))
    os.remove(image)
    return doc


def cell_text(value):
    text = character_value(value)
    return "NA" if is_missing(text) else text


def column_weights(frame):
    """How much text each column holds, so a slide table's columns share the width in proportion."""
    weights = []
    for column in frame.columns:
        longest = max((len(cell_text(value)) for value in frame[column]), default=0)
        weights.append(max(len(str(column)), longest, 3))
    return weights


def booktabs_rules(table):
    """Rules above and below the header and below the last row, as a booktabs table has."""
    from pptx.oxml.ns import qn

    last = len(table.rows) - 1
    for column in range(len(table.columns)):
        for row, edges in ((0, ("a:lnT", "a:lnB")), (last, ("a:lnB",))):
            cell_properties = table.cell(row, column)._tc.get_or_add_tcPr()
            for edge in edges:
                line = cell_properties.makeelement(qn(edge), {"w": "12700"})
                fill = line.makeelement(qn("a:solidFill"), {})
                fill.append(fill.makeelement(qn("a:srgbClr"), {"val": "000000"}))
                line.append(fill)
                cell_properties.append(line)


def style_cell(cell, text, font_size, bold):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    cell.text = text
    cell.fill.background()
    for edge in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(cell, edge, Pt(TABLE_PADDING_POINTS))
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.size = Pt(font_size)
    run.font.bold = bold


def slide_table(doc, frame, font_size):
    """The table on the current slide, spanning the content placeholder and styled like R's booktabs flextable."""
    from pptx.util import Inches

    slot = content_slot(doc)
    height = (len(frame) + 1) * font_size * 2.0 / 72
    shape = current_slide(doc).shapes.add_table(len(frame) + 1, len(frame.columns), Inches(slot["left"]), Inches(slot["top"]), Inches(slot["width"]), Inches(height))
    table = shape.table
    shape._element.graphic.graphicData.tbl.tblPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId"
    ).text = NO_GRID_STYLE
    weights = column_weights(frame)
    for i, weight in enumerate(weights):
        table.columns[i].width = Inches(slot["width"] * weight / sum(weights))
    for j, column in enumerate(frame.columns):
        style_cell(table.cell(0, j), str(column), font_size, True)
        for i, value in enumerate(frame[column]):
            style_cell(table.cell(i + 1, j), cell_text(value), font_size, False)
    booktabs_rules(table)
    return doc


def word_table(doc, frame, font_size, autofit):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    table = doc.add_table(rows=len(frame) + 1, cols=len(frame.columns))
    table.autofit = bool(autofit)
    rows = [[str(column) for column in frame.columns]]
    for record in frame.itertuples(index=False):
        rows.append([cell_text(value) for value in record])
    for i, values in enumerate(rows):
        for j, text in enumerate(values):
            cell = table.cell(i, j)
            cell.text = text
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                run.font.bold = i == 0
    return doc


def report_add_table(doc, data, font_size=TABLE_FONT_SIZE, autofit=True):
    """Add a table to a report.

    Adds a data frame in the current slide's content placeholder (pptx) or as a
    new table (docx). On a slide the table is styled (rules above and below the
    header, centred cells, a readable font) and its columns are widened to span
    the content placeholder, roughly in proportion to each column's contents,
    so it fills the slide instead of sitting tiny in a corner. In Word it is
    added as a plain table with a bold header.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    data : pandas.DataFrame
        The table.
    font_size : float
        Cell font size in points. Default 12.
    autofit : bool
        Word only: auto-size columns to content. Default True.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_add_plot

    Examples
    --------
    >>> doc = report_add_table(report_new("docx"), calc_percentage(podracing_survey, "demo_gender"))
    """
    check_doc(doc)
    frame = pd.DataFrame(data).reset_index(drop=True)
    if is_pptx(doc):
        return slide_table(doc, frame, font_size)
    return word_table(doc, frame, font_size, autofit)


def report_add_text(doc, text, **kwargs):
    """Add a paragraph of text to a report.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    text : str or list of str
        Text to add; a list becomes one paragraph (bullet) per element on
        slides.
    **kwargs
        Passed to python-docx's ``add_paragraph()`` (docx only).

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_slide

    Examples
    --------
    >>> doc = report_add_text(report_new("docx"), "Key findings follow.")
    """
    check_doc(doc)
    lines = [text] if isinstance(text, str) else list(text)
    if not is_pptx(doc):
        for line in lines:
            doc.add_paragraph(line, **kwargs)
        return doc
    state = report_state(doc)
    bodies = []
    if state.get("layout") is not None:
        bodies = layout_placeholders(doc, state["layout"], state["master"], "body")
    if bodies:
        fill_placeholder(doc, bodies[0], lines)
        return doc
    from pptx.util import Inches

    slot = content_slot(doc)
    box = current_slide(doc).shapes.add_textbox(Inches(slot["left"]), Inches(slot["top"]), Inches(slot["width"]), Inches(slot["height"]))
    box.text_frame.text = lines[0]
    for line in lines[1:]:
        box.text_frame.add_paragraph().text = line
    return doc


def report_slide(doc, title=None, content=None, layout=None, master=None, **kwargs):  # lint-style: ignore FN003
    """Add a whole slide in one call (title plus its content).

    The one-line-per-slide wrapper: starts a new slide, sets its title, and
    places `content` on it, so a deck script reads as one line per slide.
    `content` is dispatched by type: a chart becomes an image, a data frame a
    table, text a bulleted text box. Equivalent to `report_add_slide()`
    followed by the matching ``report_add_*()`` call.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    title : str, optional
        Slide title, typically the survey question the slide answers.
    content : plotnine.ggplot, pandas.DataFrame, str or list of str, optional
        What the slide shows. None (default) adds an empty titled slide.
    layout, master : str, optional
        Passed to `report_add_slide()`; None auto-selects.
    **kwargs
        Passed on to `report_add_plot()`, `report_add_table()` or
        `report_add_text()` depending on `content`.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_section, report_title_slide, report_add_slide

    Examples
    --------
    >>> chart = plot_bars(calc_percentage(podracing_survey, "demo_gender"))
    >>> doc = report_slide(report_new("pptx"), "Who follows pod racing?", chart)
    """
    check_doc(doc)
    doc = report_add_slide(doc, title=title, layout=layout, master=master)
    if content is None:
        return doc
    if is_ggplot(content):
        return report_add_plot(doc, content, **kwargs)
    if isinstance(content, pd.DataFrame):
        return report_add_table(doc, content, **kwargs)
    if isinstance(content, (str, list, tuple)):
        return report_add_text(doc, content, **kwargs)
    raise ValueError("`content` must be a ggplot, a data frame, or a character vector.")


def report_section(doc, title, layout="Section Header", master=None):
    """Add a section-divider slide.

    A one-line section break: a slide on the template's "Section Header"
    layout, falling back to a plain titled slide when the template has none.
    Use a short, single-word label ("DEMOGRAPHICS", "RATINGS", "APPENDIX") to
    chapter a deck.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    title : str
        Section label.
    layout : str
        Section layout name. Default "Section Header".
    master : str, optional
        Master name.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_slide, report_title_slide

    Examples
    --------
    >>> doc = report_section(report_new("pptx"), "DEMOGRAPHICS")
    """
    check_doc(doc)
    if not is_pptx(doc):
        return report_add_slide(doc, title, heading_level=1)
    if layout in report_layouts(doc)["layout"].tolist():
        return report_add_slide(doc, title, layout=layout, master=master)
    return report_add_slide(doc, title)


def report_title_slide(doc, title, subtitle=None, layout=None, master=None):
    """Add the opening title slide.

    A one-line title slide, placed on the template's title layout (the one with
    a centre-title placeholder), with an optional strapline in its subtitle
    placeholder.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    title : str
        Deck title.
    subtitle : str, optional
        Strapline for the subtitle placeholder: the respondent base and
        fieldwork period, say.
    layout, master : str, optional
        Overrides; None auto-selects the title layout.

    Returns
    -------
    pptx.Presentation or docx.Document
        The document.

    See Also
    --------
    report_slide, report_section

    Examples
    --------
    >>> doc = report_title_slide(report_new("pptx"), "Pod-Racing Fan Survey", subtitle="1,000 fans | Fieldwork 2026")
    """
    check_doc(doc)
    if not is_pptx(doc):
        doc = report_add_slide(doc, title, heading_level=1)
        if subtitle is not None:
            doc.add_paragraph(subtitle)
        return doc
    if layout is None:
        selection = select_layout(doc, "title")
        layout = selection["layout"]
        master = master or selection["master"]
    doc = report_add_slide(doc, title, layout=layout, master=master)
    if subtitle is not None:
        subtitles = layout_placeholders(doc, layout, master or report_state(doc)["master"], ["subTitle", "body"])
        if subtitles:
            fill_placeholder(doc, subtitles[0], subtitle)
    return doc


def report_save(doc, path=None):
    """Save a report to disk.

    Writes a deck or document built with `report_new()` and the ``report_*``
    verbs, and returns the path it used. The format follows the document, not
    the file name: a deck writes .pptx and a Word document writes .docx, which
    is why `path` may be left out entirely.

    Parameters
    ----------
    doc : pptx.Presentation or docx.Document
        A document from `report_new()`.
    path : str, optional
        Output file path (.pptx or .docx). None (default) writes
        ``report.pptx`` / ``report.docx``. A bare file name lands in
        ``ezrsurvey-outputs/`` (created on demand); a path naming a directory
        ("./x.pptx", "charts/x.pptx", anything absolute) is used exactly as
        given. See the ``output_dir`` option.

    Returns
    -------
    str
        The path written.

    See Also
    --------
    report_new, report_deck

    Examples
    --------
    >>> import tempfile, os
    >>> path = os.path.join(tempfile.mkdtemp(), "deck.pptx")
    >>> report_save(report_add_slide(report_new("pptx"), "Hi"), path)
    """
    check_doc(doc)
    ext = "pptx" if is_pptx(doc) else "docx"
    path = resolve_output_path(path if path is not None else default_output_path("report", ext))
    doc.save(path)
    return path
