"""Port R's test-report.R: Quarto scaffolds, layout selection, slide and Word builders, slide numbers and the example."""

import os
import zipfile

import pytest
from pptx import Presentation

import ezrsurvey as ez
from ezrsurvey.plot import consistent_bar_width
from ezrsurvey.report_builder import TEMPLATE_FOLDER, select_layout


def gender_bars():
    return ez.plot_bars(ez.calc_percentage(ez.podracing_survey, "demo_gender"))


def slide_xml(path, number):
    with zipfile.ZipFile(path) as archive:
        return archive.read(f"ppt/slides/slide{number}.xml").decode("utf-8")


def test_scaffold_report_writes_the_title(tmp_path):
    path = str(tmp_path / "r.qmd")
    assert ez.scaffold_report("html", path=path, title="My Survey") == path
    text = open(path, encoding="utf-8").read()
    assert "My Survey" in text
    assert "format:" in text
    with pytest.raises(ValueError):
        ez.scaffold_report("html", path=path)


def test_scaffold_report_substitutes_author_and_reference_doc(tmp_path):
    path = str(tmp_path / "r.qmd")
    reference = tmp_path / "brand.pptx"
    reference.write_bytes(b"")
    ez.scaffold_report("pptx", path=path, title="T", author="A. Analyst", reference_doc=str(reference))
    text = open(path, encoding="utf-8").read()
    assert "A. Analyst" in text
    assert "reference-doc:" in text
    assert "{{" not in text
    second = str(tmp_path / "r2.qmd")
    ez.scaffold_report("pptx", path=second, title="T")
    assert "ezrsurvey-16x9.pptx" in open(second, encoding="utf-8").read()
    third = str(tmp_path / "r3.qmd")
    ez.scaffold_report("docx", path=third, title="T")
    assert "# reference-doc:" in open(third, encoding="utf-8").read()


def test_bundled_templates_use_python_and_parameters():
    for path in TEMPLATE_FOLDER.glob("report-*.qmd"):
        text = path.read_text(encoding="utf-8")
        assert "tags: [parameters]" in text, path.name
        assert "```{r" not in text, path.name


def test_list_report_templates_lists_the_formats():
    assert {"pptx", "html", "pdf", "docx"} <= set(ez.list_report_templates())


def test_report_layouts_describes_the_default_template():
    layouts = ez.report_layouts()
    assert {"layout", "master", "has_title", "n_body"} <= set(layouts.columns)
    assert "Title and Content" in layouts["layout"].tolist()


def test_select_layout_picks_by_placeholder_types():
    doc = ez.report_new("pptx")
    assert select_layout(doc, "content")["layout"] == "Title and Content"
    assert select_layout(doc, "title")["layout"] == "Title Slide"
    two = select_layout(doc, "two_content")
    assert two is None or two["layout"] == "Two Content"


def test_report_add_slide_validates_layouts_and_guards_titles():
    doc = ez.report_new("pptx")
    with pytest.raises(ValueError, match="report_layouts"):
        ez.report_add_slide(doc, layout="No Such Layout")
    with pytest.warns(UserWarning, match="no title placeholder"):
        ez.report_add_slide(doc, title="Hi", layout="Blank")


def test_report_deck_builds_a_pptx(tmp_path):
    path = str(tmp_path / "deck.pptx")
    table = ez.calc_nps(ez.podracing_survey, "nps_value")
    out = ez.report_deck({"Gender": gender_bars(), "NPS": table}, path=path, title="Overview")
    assert os.path.getsize(out) > 0
    assert len(Presentation(out).slides) == 3


def test_word_reports_take_plots_and_tables(tmp_path):
    chart = ez.plot_bars(ez.calc_percentage(ez.podracing_survey, "demo_edu"), flip=True)
    doc = ez.report_new("docx")
    doc = ez.report_add_slide(doc, "Education", heading_level=1)
    doc = ez.report_add_plot(doc, chart)
    doc = ez.report_add_table(doc, ez.calc_percentage(ez.podracing_survey, "demo_edu"))
    out = ez.report_save(doc, str(tmp_path / "doc.docx"))
    assert os.path.exists(out)


def test_the_default_deck_is_16_by_9():
    doc = ez.report_new("pptx")
    assert doc.slide_width / 914400 == pytest.approx(40 / 3, abs=1e-6)
    assert doc.slide_height / 914400 == pytest.approx(7.5)
    layouts = ez.report_layouts()
    assert {"Title Slide", "Title and Content", "Two Content", "Content with Caption"} <= set(layouts["layout"])
    assert layouts.loc[layouts["layout"] == "Title and Content", "n_body"].tolist() == [1]


def test_content_slides_carry_a_slide_number(tmp_path):
    doc = ez.report_new("pptx")
    doc = ez.report_add_slide(doc, "One")
    doc = ez.report_add_slide(doc, "Two")
    path = ez.report_save(doc, str(tmp_path / "numbered.pptx"))
    for number in (1, 2):
        assert "slidenum" in slide_xml(path, number)
    plain = ez.report_new("pptx", slide_numbers=False)
    plain = ez.report_add_slide(plain, "One")
    path = ez.report_save(plain, str(tmp_path / "plain.pptx"))
    assert "slidenum" not in slide_xml(path, 1)


def test_bars_keep_a_constant_thickness():
    assert consistent_bar_width(3) / 3 == pytest.approx(consistent_bar_width(8) / 8)
    assert consistent_bar_width(3) < consistent_bar_width(8)
    assert consistent_bar_width(40) <= 0.9
    assert consistent_bar_width(6) == pytest.approx(0.66)


def test_report_deck_builds_a_multi_chart_deck(tmp_path):
    items = {"Gender": gender_bars(), "Job": ez.plot_bars(ez.calc_percentage(ez.podracing_survey, "demo_job"))}
    assert os.path.exists(ez.report_deck(items, path=str(tmp_path / "deck.pptx")))


def test_one_line_slide_verbs_build_a_deck(tmp_path):
    doc = ez.report_new("pptx")
    doc = ez.report_title_slide(doc, "Pod-Racing Fan Survey")
    doc = ez.report_section(doc, "DEMOGRAPHICS")
    doc = ez.report_slide(doc, "What is your gender?", gender_bars())
    doc = ez.report_slide(doc, "Where from?", ez.calc_nps(ez.podracing_survey, "nps_value", by="region"))
    doc = ez.report_slide(doc, "How to read this", ["First point", "Second point"])
    out = ez.report_save(doc, str(tmp_path / "wrappers.pptx"))
    assert os.path.getsize(out) > 0


def test_report_title_slide_places_a_subtitle(tmp_path):
    doc = ez.report_title_slide(ez.report_new("pptx"), "Deck", subtitle="1,000 fans | Fieldwork 2026")
    path = ez.report_save(doc, str(tmp_path / "cover.pptx"))
    assert "Fieldwork 2026" in slide_xml(path, 1)


def test_report_section_falls_back_when_the_layout_is_absent():
    doc = ez.report_new("pptx")
    ez.report_section(doc, "RATINGS")
    ez.report_section(doc, "NOPE", layout="No Such Layout")
    assert len(doc.slides) == 2


def test_report_slide_rejects_unsupported_content():
    with pytest.raises(ValueError, match="ggplot"):
        ez.report_slide(ez.report_new("pptx"), "T", content=42)


def test_example_report_copies_the_worked_example(tmp_path):
    destination = str(tmp_path / "example")
    paths = ez.example_report(destination)
    assert all(os.path.exists(path) for path in paths)
    assert any(path.endswith("podracing-report.qmd") for path in paths)
    assert any(path.endswith("podracing-deck.py") for path in paths)
    with pytest.raises(ValueError, match="overwrite"):
        ez.example_report(destination)
    ez.example_report(destination, overwrite=True)


def test_report_new_can_start_from_an_empty_deck():
    kept = ez.report_new("pptx")
    emptied = ez.report_new("pptx", keep_slides=False)
    assert len(emptied.slides) == 0
    assert len(kept.slides) >= len(emptied.slides)
